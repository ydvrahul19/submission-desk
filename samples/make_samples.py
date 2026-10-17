"""Generates three fictional submissions (.eml with PDF/DOCX/XLSX/CSV attachments) into samples/<name>/.
Dates are relative to the day you run it:  python samples/make_samples.py"""
import io
import sys
from datetime import date, timedelta
from email.message import EmailMessage
from pathlib import Path


def pdf_bytes(pages):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for lines in pages:
        y = 740
        for ln in lines:
            c.drawString(60, y, ln)
            y -= 20
        c.showPage()
    c.save()
    return buf.getvalue()


def xlsx_bytes(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SOV"
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def docx_bytes(lines):
    import docx
    d = docx.Document()
    for ln in lines:
        d.add_paragraph(ln)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def build(root: Path, today: date):
    iso = lambda n: (today + timedelta(days=n)).isoformat()
    cases = {
        "1_harbor_bakes_clean": dict(
            subject="New submission - Harbor Bakes LLC", body=f"Hi team, please quote the attached bakery.\nNamed insured: Harbor Bakes LLC\nEffective date: {iso(35)}",
            atts=[("application.pdf", "application/pdf", pdf_bytes([["APPLICATION - PROPERTY", "Insured: Harbor Bakes LLC", "Operations: Retail bakery and cafe", "Year built: 2008", "Sprinklers: yes", "Losses in last 5 years: 1"]])),
                  ("sov.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_bytes([["Statement of values"], ["Total insured value", "$3,400,000"], ["Location", "14 Pier Road"]]))]),
        "2_redline_pyrotech_outside_appetite": dict(
            subject="Submission - Redline Pyrotech", body=f"Renewal details attached.\nApplicant: Redline Pyrotech Inc\nEffective date: {iso(40)}",
            atts=[("application.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes(["APPLICATION - PROPERTY", "Insured: Redline Pyrotech Inc", "Operations: Fireworks storage and distribution", "Year built: 1974", "Sprinklers: no", "Claims in past 5 years: 3"])),
                  ("sov.csv", "text/csv", b"Statement of values\nTIV,$14.0 million\n")]),
        "3_meridian_cold_storage_conflicts": dict(
            subject="Meridian Cold Storage - rush", body=f"Need terms fast.\nNamed insured: Meridian Cold Storage Ltd\nTotal insured value: $8,200,000\nEffective date: {iso(6)}",
            atts=[("application.pdf", "application/pdf", pdf_bytes([["APPLICATION - PROPERTY", "Insured: Meridian Cold Storage Ltd"], ["Operations: Refrigerated warehouse", "Year built: 1991", "Sprinklers: partial"]])),
                  ("sov.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_bytes([["Statement of values"], ["Total insured value", "$12,500,000"]]))]),
    }
    for name, c in cases.items():
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = "broker@example.com", "submissions@example.com", c["subject"]
        m.set_content(c["body"])
        for fn, mime, data in c["atts"]:
            mt, st = mime.split("/")
            m.add_attachment(data, maintype=mt, subtype=st, filename=fn)
        d = root / name
        d.mkdir(parents=True, exist_ok=True)
        (d / "submission.eml").write_bytes(bytes(m))


if __name__ == "__main__":
    build(Path(__file__).parent, date.today())
    print("samples written")
