"""Read an email and attachments (.eml .pdf .docx .xlsx .csv .txt) into Doc objects."""
import csv
import io
from email import policy
from email.parser import BytesParser
from typing import List

from .models import Doc


def read_pdf(name: str, data: bytes) -> Doc:
    from pypdf import PdfReader
    pages = [(p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages]
    return Doc(name, pages if any(x.strip() for x in pages) else ["[no extractable text - scanned PDF? OCR is not included]"])


def read_docx(name: str, data: bytes) -> Doc:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(f"{cells[0]}: {' | '.join(cells[1:])}" if len(cells) > 1 else cells[0])
    return Doc(name, ["\n".join(lines)])


def read_xlsx(name: str, data: bytes) -> Doc:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    pages = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = [str(c) for c in r if c is not None]
            if len(cells) > 1:
                rows.append(f"{cells[0]}: {' | '.join(cells[1:])}")
            elif cells:
                rows.append(cells[0])
        pages.append(f"[sheet {ws.title}]\n" + "\n".join(rows))
    return Doc(name, pages)


def read_csv(name: str, data: bytes) -> Doc:
    text = data.decode("utf-8", errors="replace")
    rows = [f"{r[0]}: {' | '.join(r[1:])}" if len(r) > 1 else (r[0] if r else "") for r in csv.reader(io.StringIO(text))]
    return Doc(name, ["\n".join(rows)])


def read_eml(name: str, data: bytes) -> List[Doc]:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    head = f"From: {msg['from']}\nSubject: {msg['subject']}\n"
    body = msg.get_body(preferencelist=("plain", "html"))
    docs = [Doc(f"{name} (email body)", [head + (body.get_content() if body else "")])]
    for part in msg.iter_attachments():
        docs += read_file(part.get_filename() or "attachment", part.get_payload(decode=True) or b"")
    return docs


def read_file(name: str, data: bytes) -> List[Doc]:
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    try:
        if ext == "eml":
            return read_eml(name, data)
        if ext == "pdf":
            return [read_pdf(name, data)]
        if ext == "docx":
            return [read_docx(name, data)]
        if ext in ("xlsx", "xlsm"):
            return [read_xlsx(name, data)]
        if ext == "csv":
            return [read_csv(name, data)]
        return [Doc(name, [data.decode("utf-8", errors="replace")])]
    except Exception as e:  # a corrupt or protected file must not crash the run
        return [Doc(name, [f"[could not read this file: {type(e).__name__}]"])]
